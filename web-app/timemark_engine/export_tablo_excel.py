#!/usr/bin/env python3
"""
scripts/export_tablo_excel.py
------------------------------
Generates official KAI Form No. STE-RECORD-13.4.01:
"JADWAL CHECKLIST & PERAWATAN BERKALA (UPT RESORT SINTEL 1.21 BOGOR)".

Modes:
- Mode 1 (pipeline): Uses active schedule.json
- Mode 2 (custom): Runs scheduler.py on the specified custom PDF directory
"""

import sys
import os

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

import re
import json
import calendar
import subprocess
from pathlib import Path
from datetime import datetime, date
import argparse

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

APP_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(APP_DIR))
sys.path.insert(0, str(APP_DIR / "scripts"))

try:
    from employee_manager import (
        load_pegawai_config,
        extract_page1_tim1_personnel,
        get_tim2_roster_for_date
    )
except ImportError:
    pass

import fitz

MONTH_NAMES_ID = {
    1: "JANUARI", 2: "FEBRUARI", 3: "MARET", 4: "APRIL",
    5: "MEI", 6: "JUNI", 7: "JULI", 8: "AGUSTUS",
    9: "SEPTEMBER", 10: "OKTOBER", 11: "NOVEMBER", 12: "DESEMBER"
}

DAY_NAMES_ID = {
    0: "Senin", 1: "Selasa", 2: "Rabu", 3: "Kamis",
    4: "Jumat", 5: "Sabtu", 6: "Minggu"
}

# 46 Equipment Groups matching official Tablo layout
TABLO_GROUPS = {
    # 1. BOGOR (1 - 8)
    1: {"station": "BOGOR", "name": "Wesel Utara"},
    2: {"station": "BOGOR", "name": "Wesel Selatan"},
    3: {"station": "BOGOR", "name": "PDSE, VDU, CTS, UPS, Genset, FO, Telkom WS 3bln"},
    4: {"station": "BOGOR", "name": "AxleCounter Siemens Utara"},
    5: {"station": "BOGOR", "name": "AxleCounter Siemens Selatan"},
    6: {"station": "BOGOR", "name": "Peraga sinyal area Utara"},
    7: {"station": "BOGOR", "name": "Peraga sinyal area Selatan"},
    8: {"station": "BOGOR", "name": "K3, Radio, FO, UPS & Genset, WS 3BLN, JPL 01 & 02"},

    # 2. CILEBUT - BOJONGGEDE (9 - 10)
    9: {"station": "CILEBUT-BOJONGGEDE", "name": "Sinyal Blok Otomatis"},
    10: {"station": "CILEBUT-BOJONGGEDE", "name": "AxleCounter Siemens"},

    # 3. BOP / PALEDANG (11 - 16)
    11: {"station": "BOP", "name": "Wesel"},
    12: {"station": "BOP", "name": "AxleCounter Frauscher Sec 1"},
    13: {"station": "BOP", "name": "Peraga Sinyal Sec 1"},
    14: {"station": "BOP", "name": "AxleCounter Frauscher Sec 2"},
    15: {"station": "BOP", "name": "Peraga Sinyal Sec 2"},
    16: {"station": "BOP", "name": "PDSE, VDU, Catu daya, Genset, FO, Telkom, JPL 04, WS 3bln"},

    # 4. BOGOR - CILEBUT (17 - 25)
    17: {"station": "BOGOR-CILEBUT", "name": "Peraga Sinyal Blok 1"},
    18: {"station": "BOGOR-CILEBUT", "name": "Peraga Sinyal Blok 2"},
    19: {"station": "BOGOR-CILEBUT", "name": "Peraga Sinyal Blok 3"},
    20: {"station": "BOGOR-CILEBUT", "name": "Peraga Sinyal Blok 4"},
    21: {"station": "BOGOR-CILEBUT", "name": "AxleCounter Siemens 1"},
    22: {"station": "BOGOR-CILEBUT", "name": "AxleCounter Siemens 2"},
    23: {"station": "BOGOR-CILEBUT", "name": "AxleCounter Siemens 3"},
    24: {"station": "BOGOR-CILEBUT", "name": "AxleCounter Siemens 4"},
    25: {"station": "BOGOR-CILEBUT", "name": "Pintu Perlintasan JPL 27 , JPL 28"},

    # 5. CILEBUT (26 - 30)
    26: {"station": "CILEBUT", "name": "Wesel"},
    27: {"station": "CILEBUT", "name": "Peraga Sinyal"},
    28: {"station": "CILEBUT", "name": "AxleCounter Siemens"},
    29: {"station": "CILEBUT", "name": "PDSE, CTDYA, Telkom, VDU, FO, WS 3bln"},
    30: {"station": "CILEBUT", "name": "Pintu perlintasan JPL 26N"},

    # 6. BATUTULIS (31 - 34)
    31: {"station": "BATUTULIS", "name": "Peraga Sinyal"},
    32: {"station": "BATUTULIS", "name": "AxleCounter Frauscher"},
    33: {"station": "BATUTULIS", "name": "PDSE, VDU, Genset, Battery Bank"},
    34: {"station": "BATUTULIS", "name": "FO, Telkom, Catu daya, K3 ER, WS 3bln"},

    # 7. MASENG (35 - 38)
    35: {"station": "MASENG", "name": "Wesel"},
    36: {"station": "MASENG", "name": "AxleCounter Frauscher"},
    37: {"station": "MASENG", "name": "Peraga Sinyal"},
    38: {"station": "MASENG", "name": "PDSE, VDU, Genset, FO, PTLS, Catu daya, K3 ER, WS 3bln"},

    # 8. PINTU PERLINTASAN BOO-BTT (39)
    39: {"station": "PINTU PERLINTASAN BOO-BTT", "name": "Pintu Perlintasan JPLE 11 , JPLE BNR, JPLE 07"},

    # 9. IB CIOMAS (40 - 41)
    40: {"station": "IB CIOMAS", "name": "Peraga Sinyal, AxleCounter Frausher"},
    41: {"station": "IB CIOMAS", "name": "PDSE, PTLS, Catudaya, FO, Tower, K3 ER, Basestation 6bln"},

    # 10. CIGOMBONG (42 - 46)
    42: {"station": "CIGOMBONG", "name": "Peraga Sinyal"},
    43: {"station": "CIGOMBONG", "name": "Axle counter Frausher"},
    44: {"station": "CIGOMBONG", "name": "PDSE, PTLS, Catudaya, FO, Genset, K3 ER"},
    45: {"station": "CIGOMBONG", "name": "Jpl 15, Jpl 16"},
    46: {"station": "CIGOMBONG", "name": "Basestation 6 bulanan"},
}

STATION_COLORS = {
    "BOGOR": "FFE599",                 # Yellow gold
    "CILEBUT-BOJONGGEDE": "D9EAD3",    # Soft green
    "BOP": "CFE2F3",                   # Soft blue
    "BOGOR-CILEBUT": "FFF2CC",         # Pale yellow
    "CILEBUT": "EAD1DC",               # Soft magenta/pink
    "BATUTULIS": "FCE5CD",             # Soft orange/peach
    "MASENG": "D0E0E3",                # Soft cyan/teal
    "PINTU PERLINTASAN BOO-BTT": "F9CB9C", # Light amber
    "IB CIOMAS": "D9D2E9",             # Soft lavender
    "CIGOMBONG": "C9DAF8",             # Soft blue
}

# The 3-column grouping for the legend layout matching the photo
LEGEND_COLUMNS = [
    # Column 1
    [
        ("BOGOR", [1, 2, 3, 4, 5, 6, 7, 8]),
        ("CILEBUT-BOJONGGEDE", [9, 10]),
        ("BOP", [11, 12, 13, 14, 15, 16]),
    ],
    # Column 2
    [
        ("BOGOR-CILEBUT", [17, 18, 19, 20, 21, 22, 23, 24, 25]),
        ("CILEBUT", [26, 27, 28, 29, 30]),
        ("BATUTULIS", [31, 32, 33, 34]),
    ],
    # Column 3
    [
        ("MASENG", [35, 36, 37, 38]),
        ("PINTU PERLINTASAN BOO-BTT", [39]),
        ("IB CIOMAS", [40, 41]),
        ("CIGOMBONG", [42, 43, 44, 45, 46]),
    ]
]


def match_asset_to_tablo_group(item: dict) -> int:
    """Accurately maps any schedule asset record to one of the 46 official Tablo groups."""
    ident = (item.get("identifier") or "").upper()
    cat = (item.get("category") or "").upper()
    f = (item.get("file") or "").upper()
    text = f"{cat} {ident} {f}"

    is_boo = "BOO" in ident or "BOGOR" in f
    is_clt = "CLT" in ident or "CILEBUT" in f
    is_bjd = "BJD" in ident or "BOJONG" in f
    is_bop = "BOP" in ident or "PALEDANG" in f
    is_btt = "BTT" in ident or "BATUTULIS" in f or "BATU TULIS" in f
    is_msg = "MSG" in ident or "MASENG" in f
    is_cos = "COS" in ident or "CIOMAS" in f or " CS " in f
    is_cgb = "CGB" in ident or "CIGOMBONG" in f

    is_boo_clt = ("BOO-CLT" in text or "CLT-BOO" in text or ("BOO" in text and "CLT" in text and "JPL" in text))
    is_bjd_clt = "BJD-CLT" in text or "CLT-BJD" in text
    is_boo_btt = "BOO-BTT" in text or "BTT-BOO" in text or "BOP-BTT" in text or "BOO-BOP" in text or "BTT-BOP" in text

    # PINTU PERLINTASAN BOO-BTT (39)
    if is_boo_btt and ("JPL" in text or "PTPP" in text):
        return 39

    # CIGOMBONG (42 - 46)
    if is_cgb:
        if "RADIO" in text or "BASESTATION" in text: return 46
        if "JPL" in text or "PTPP" in text: return 45
        if cat in ["PDSE", "PTDS", "PTLS", "CATUDAYA", "SERAT OPTIK"]: return 44
        if cat == "AXC": return 43
        if cat == "SINYAL": return 42
        return 44

    # IB CIOMAS (40 - 41)
    if is_cos:
        if cat in ["SINYAL", "AXC"]: return 40
        return 41

    # MASENG (35 - 38)
    if is_msg:
        if cat == "WESEL": return 35
        if cat == "AXC": return 36
        if cat == "SINYAL": return 37
        return 38

    # BATUTULIS (31 - 34)
    if is_btt and not is_bop and not is_boo:
        if cat == "SINYAL": return 31
        if cat == "AXC": return 32
        if cat in ["PDSE", "PTDS"]: return 33
        return 34

    # BOP (11 - 16)
    if is_bop or ("BTT" in text and "BOP" in text):
        if cat == "WESEL": return 11
        if cat == "AXC":
            if "SEC 2" in text or "BTT" in text: return 14
            return 12
        if cat == "SINYAL":
            if "SEC 2" in text or "BTT" in text or "MJ" in ident: return 15
            return 13
        return 16

    # BOGOR-CILEBUT (17 - 25)
    if is_boo_clt:
        if "JPL" in text or "PTPP" in text: return 25
        if cat == "AXC": return 21
        if cat == "SINYAL": return 17
        return 21

    # CILEBUT-BOJONGGEDE (9 - 10)
    if is_bjd_clt or (is_bjd and not is_boo):
        if cat == "SINYAL": return 9
        if cat == "AXC": return 10
        return 9

    # CILEBUT (26 - 30)
    if is_clt and not is_boo:
        if "JPL 26" in text: return 30
        if cat == "WESEL": return 26
        if cat == "SINYAL": return 27
        if cat == "AXC": return 28
        return 29

    # BOGOR (1 - 8)
    if is_boo:
        if cat == "WESEL":
            m = re.search(r"W(\d+)", ident)
            if m:
                wnum = int(m.group(1))
                if wnum in [41, 43, 61, 81]: return 2
                return 1
            return 1
        if cat in ["PDSE", "CTS", "CATUDAYA", "SERAT OPTIK", "PTLS", "PTDS", "TELEKOMUNIKASI"]:
            return 3
        if cat == "AXC":
            if any(k in ident for k in ["ZP 41", "ZP 43", "ZP 61", "ZP 81", "SELATAN"]):
                return 5
            return 4
        if cat == "SINYAL":
            if any(k in ident for k in ["SELATAN", "S1", "S2"]):
                return 7
            return 6
        if "JPL" in text or "K3" in text:
            return 8
        return 3

    return 0


def generate_tablo_workbook(schedule_data: dict, year: int = None, month: int = None, output_path: Path = None, config: dict = None, config_path: Path = None, pdf_dir: Path | str = None) -> str:
    schedules = schedule_data.get("schedules", [])

    # 1. Detect Year and Month if not provided
    if not year or not month:
        date_counts = {}
        for s in schedules:
            iso = s.get("iso_date")
            if iso and len(iso) >= 7:
                try:
                    dt = datetime.strptime(iso[:10], "%Y-%m-%d")
                    k = (dt.year, dt.month)
                    date_counts[k] = date_counts.get(k, 0) + 1
                except Exception:
                    pass
        if date_counts:
            year, month = sorted(date_counts.items(), key=lambda x: x[1], reverse=True)[0][0]
        else:
            now = datetime.now()
            year, month = now.year, now.month

    num_days = calendar.monthrange(year, month)[1]

    # 2. Load Personnel Config
    if config and isinstance(config, dict):
        cfg = config
    elif config_path and Path(config_path).exists():
        cfg = load_pegawai_config(config_path)
    else:
        cfg = load_pegawai_config()

    resor = cfg.get("resor", {})
    resor_nama = resor.get("nama", "S. SLAMET RIYADI").strip().upper()
    resor_nipp = str(resor.get("nipp", "-")).strip()
    resor_jabatan = resor.get("jabatan", "KUPT Resor Stl 1.21 Boo")
    resor_domisili = resor.get("domisili", "Bogor")

    kaur_list = cfg.get("kaur", [])
    pnc_list = cfg.get("pnc", [])

    personnel_list = []
    # 1. Resor
    personnel_list.append({
        "key": "resor",
        "nama": resor_nama,
        "full_name": resor_nama,
        "nipp": resor_nipp,
        "jabatan": resor_jabatan,
        "domisili": resor_domisili,
        "is_resor": True
    })

    # 2. KAUR
    for idx, k in enumerate(kaur_list):
        k_name = k.get("nama", "").strip().upper()
        k_key = f"kaur_{idx}_{re.sub(r'[^a-zA-Z0-9]', '', k_name.lower())}"
        personnel_list.append({
            "key": k_key,
            "nama": k_name,
            "full_name": k_name,
            "nipp": str(k.get("nipp", "-")).strip(),
            "jabatan": k.get("jabatan", "Kepala Urusan Preventif" if idx == 0 else "Kepala Urusan Perbaikan"),
            "domisili": k.get("domisili", "Bogor"),
            "role": "KAUR"
        })

    # 3. PNC
    for idx, p in enumerate(pnc_list):
        p_name = p.get("nama", "").strip().upper()
        p_key = f"pnc_{idx}_{re.sub(r'[^a-zA-Z0-9]', '', p_name.lower())}"
        personnel_list.append({
            "key": p_key,
            "nama": p_name,
            "full_name": p_name,
            "nipp": str(p.get("nipp", "-")).strip(),
            "jabatan": p.get("jabatan", "Petugas Negatif Check"),
            "domisili": p.get("domisili", "Bogor"),
            "role": "PNC"
        })

    num_p = len(personnel_list)
    p_start_col = 4
    p_end_col = p_start_col + num_p - 1
    gap_col = p_end_col + 1
    right_start_col = gap_col + 1
    right_end_col = right_start_col + 7

    # 3. Organize schedule by day (1..num_days) and Tim (1 or 2)
    daily_tim1_groups = {d: set() for d in range(1, num_days + 1)}
    daily_tim2_groups = {d: set() for d in range(1, num_days + 1)}
    daily_tim1_files = {d: [] for d in range(1, num_days + 1)}
    daily_tim2_files = {d: [] for d in range(1, num_days + 1)}

    for s in schedules:
        iso = s.get("iso_date")
        if not iso:
            continue
        try:
            dt = datetime.strptime(iso[:10], "%Y-%m-%d")
            if dt.year == year and dt.month == month:
                d = dt.day
                grp = match_asset_to_tablo_group(s)
                tim = s.get("tim", 1)
                f_name = s.get("file", "")
                if tim == 2:
                    if grp > 0: daily_tim2_groups[d].add(grp)
                    if f_name: daily_tim2_files[d].append(f_name)
                else:
                    if grp > 0: daily_tim1_groups[d].add(grp)
                    if f_name: daily_tim1_files[d].append(f_name)
        except Exception:
            continue

    # Pre-scan Tim 1 personnel from available PDFs in pdf_dir, 01_pdf_source or 02_pdf_target if available
    pdf_by_date = {}
    search_dirs = []
    if pdf_dir and Path(pdf_dir).exists():
        search_dirs.append(Path(pdf_dir))
    search_dirs.extend([Path("01_pdf_source"), Path("02_pdf_target")])
    for search_dir in search_dirs:
        if search_dir.exists():
            for pdf_file in search_dir.rglob("*.pdf"):
                m = re.search(r'(\d{2})-(\d{2})-(\d{4})', pdf_file.name)
                if m:
                    d_str = m.group(0)
                    if d_str not in pdf_by_date:
                        pdf_by_date[d_str] = pdf_file

    def _match_token_to_key(tok_str: str) -> str | None:
        tok = str(tok_str).strip().upper()
        clean_tok = re.sub(r'[^A-Z]', '', tok)
        if not clean_tok:
            return None
        for p in personnel_list[1:]:
            p_u = p["nama"].upper()
            clean_p = re.sub(r'[^A-Z]', '', p_u)
            if clean_p and clean_tok and (clean_p == clean_tok or (len(clean_p) >= 6 and clean_p in clean_tok) or (len(clean_tok) >= 6 and clean_tok in clean_p)):
                return p["key"]
            chars = [re.escape(c) for c in p_u if not c.isspace()]
            if chars and re.search(r'\b' + r'\s*'.join(chars) + r'\b', tok):
                return p["key"]
            if any(part in tok for part in p_u.split() if len(part) > 3):
                return p["key"]
        return None

    # Build matrix rows: day -> dict {person_key: cell_value}
    matrix_data = {}
    for day in range(1, num_days + 1):
        d_str = f"{day:02d}-{month:02d}-{year}"
        iso_str = f"{year:04d}-{month:02d}-{day:02d}"

        person_groups = {p["key"]: set() for p in personnel_list}
        all_day_groups = set()

        # 1. Kumpulkan seluruh aset pada tanggal ini dan cocokkan personilnya
        entries_today = [s for s in schedules if s.get("iso_date") == iso_str]

        for s in entries_today:
            grp = match_asset_to_tablo_group(s)
            if grp > 0:
                all_day_groups.add(grp)
                entry_tokens = []
                if s.get("personnel_override"):
                    for name in re.split(r"[,;\n]+", str(s["personnel_override"])):
                        if name.strip():
                            entry_tokens.append(name.strip().upper())
                if s.get("personnel"):
                    for p_item in s["personnel"]:
                        if str(p_item).strip():
                            entry_tokens.append(str(p_item).strip().upper())

                for tok in entry_tokens:
                    pkey = _match_token_to_key(tok)
                    if pkey:
                        person_groups[pkey].add(grp)

        # 2. Jika belum ada personil terpetakan sama sekali dari schedule, coba scan dari PDF jika ada
        any_assigned = any(len(grps) > 0 for k, grps in person_groups.items() if k != "resor")
        if (len(daily_tim1_files[day]) > 0 or len(daily_tim2_files[day]) > 0) and not any_assigned and d_str in pdf_by_date:
            try:
                doc = fitz.open(pdf_by_date[d_str])
                pdf_tokens = extract_page1_tim1_personnel(doc)
                doc.close()
                for tok in pdf_tokens:
                    pkey = _match_token_to_key(tok)
                    if pkey:
                        t1_grps = daily_tim1_groups.get(day, set())
                        person_groups[pkey].update(t1_grps or all_day_groups)
            except Exception:
                pass

        # 3. Fallback deterministic roster HANYA jika hari tersebut memiliki aset tetapi SAMA SEKALI tidak ada personil yang tercatat
        any_assigned = any(len(grps) > 0 for k, grps in person_groups.items() if k != "resor")
        if all_day_groups and not any_assigned:
            all_kaurs = [p for p in personnel_list if p.get("role") == "KAUR"]
            all_pncs = [p for p in personnel_list if p.get("role") == "PNC"]
            t1_k = set()
            t2_k = set()
            if all_kaurs:
                k_idx = (day - 1) % len(all_kaurs)
                t1_k.add(all_kaurs[k_idx]["key"])
                for k in all_kaurs:
                    if k["key"] not in t1_k:
                        t2_k.add(k["key"])
            if len(all_pncs) >= 2:
                p_shift = ((day - 1) * 2) % len(all_pncs)
                pnc_t1 = [all_pncs[p_shift % len(all_pncs)], all_pncs[(p_shift + 1) % len(all_pncs)]]
                for p in pnc_t1:
                    t1_k.add(p["key"])
                for p in all_pncs:
                    if p["key"] not in t1_k:
                        t2_k.add(p["key"])

            t1_grps = daily_tim1_groups.get(day, set())
            t2_grps = daily_tim2_groups.get(day, set())
            for k in t1_k:
                person_groups[k].update(t1_grps or all_day_groups)
            for k in t2_k:
                person_groups[k].update(t2_grps)

        day_values = {}
        # Resor (KUPT) schedule: selalu dinas jika ada perawatan, atau dinas kantor 'P'
        resor_key = personnel_list[0]["key"]
        t1_grps = daily_tim1_groups.get(day, set())
        if t1_grps:
            day_values[resor_key] = "/".join(str(g) for g in sorted(t1_grps))
        elif all_day_groups:
            day_values[resor_key] = "/".join(str(g) for g in sorted(all_day_groups))
        else:
            day_values[resor_key] = "P"

        for p in personnel_list[1:]:
            pkey = p["key"]
            if person_groups[pkey]:
                day_values[pkey] = "/".join(str(g) for g in sorted(person_groups[pkey]))
            else:
                day_values[pkey] = "P"

        matrix_data[day] = day_values

    # 4. Create Openpyxl Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Tablo {MONTH_NAMES_ID[month]} {year}"
    ws.views.sheetView[0].showGridLines = True

    # ── Page Setup: 1-Page Fit & Center ──
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25

    # ── Styling helpers ──
    font_main = "Arial"
    f_title = Font(name=font_main, size=11, bold=True, color="000000")
    f_sub = Font(name=font_main, size=9, bold=True, color="000000")
    f_form_code = Font(name=font_main, size=8, italic=True, color="000000")
    f_header = Font(name=font_main, size=8, bold=True, color="000000")
    f_cell = Font(name=font_main, size=8, color="000000")
    f_cell_bold = Font(name=font_main, size=8, bold=True, color="000000")
    f_slogan = Font(name=font_main, size=8, italic=True, color="000000")

    al_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    al_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    al_right = Alignment(horizontal="right", vertical="center")

    thin_black = Side(style="thin", color="000000")
    border_all = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)

    fill_yellow_hdr = PatternFill(start_color="FFE599", end_color="FFE599", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_sunday = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")

    # ── Insert Logo KAI if exists ──
    logo_file = APP_DIR / "config" / "kai_logo.png"
    if logo_file.exists():
        try:
            img = OpenpyxlImage(str(logo_file))
            img.width = 115
            img.height = 42
            ws.add_image(img, "A1")
        except Exception as e:
            print(f"[WARN] Failed to insert KAI logo: {e}")

    # ── Header Text & Form Metadata ──
    # Center title
    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=gap_col)
    c_title = ws.cell(row=2, column=3, value="JADWAL CHECKLIST & PERAWATAN BERKALA")
    c_title.font = f_title
    c_title.alignment = al_center

    # Metadata at top right above the 'dibuat oleh' box
    meta_start = right_end_col - 3
    ws.merge_cells(start_row=1, start_column=meta_start, end_row=1, end_column=right_end_col)
    c_meta1 = ws.cell(row=1, column=meta_start, value="Form No. STE-RECORD-13.4.01")
    c_meta1.font = f_form_code
    c_meta1.alignment = al_right

    ws.merge_cells(start_row=2, start_column=meta_start, end_row=2, end_column=right_end_col)
    c_meta2 = ws.cell(row=2, column=meta_start, value="UPT RESORT SINTEL 1.21 BOGOR")
    c_meta2.font = Font(name=font_main, size=8.5, bold=True, color="000000")
    c_meta2.alignment = al_right

    ws.merge_cells(start_row=3, start_column=meta_start, end_row=3, end_column=right_end_col)
    c_meta3 = ws.cell(row=3, column=meta_start, value=f"{MONTH_NAMES_ID[month]} {year}")
    c_meta3.font = Font(name=font_main, size=8.5, bold=True, color="000000")
    c_meta3.alignment = al_right

    # ── Official Table Header Bars (Row 4 — SEJAJAR!) ──
    fill_blue_hdr = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
    fill_green_hdr = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")

    # WAKTU bar across A4..C4
    ws.merge_cells("A4:C4")
    ws["A4"].value = "WAKTU"
    ws["A4"].font = f_header
    ws["A4"].alignment = al_center
    ws["A4"].fill = fill_green_hdr
    for c in range(1, 4):
        ws.cell(row=4, column=c).border = border_all
        ws.cell(row=4, column=c).fill = fill_green_hdr

    # Yellow bar above personnel columns across p_start_col..p_end_col
    ws.merge_cells(start_row=4, start_column=p_start_col, end_row=4, end_column=p_end_col)
    for c in range(p_start_col, p_end_col + 1):
        ws.cell(row=4, column=c).border = border_all
        ws.cell(row=4, column=c).fill = fill_yellow_hdr

    # "dibuat oleh :" bar across right_start_col..right_end_col
    ws.merge_cells(start_row=4, start_column=right_start_col, end_row=4, end_column=right_end_col)
    c_dibuat = ws.cell(row=4, column=right_start_col, value="dibuat oleh :")
    c_dibuat.font = Font(name=font_main, size=8.5, color="000000")
    c_dibuat.alignment = al_center
    c_dibuat.fill = fill_blue_hdr
    for c in range(right_start_col, right_end_col + 1):
        ws.cell(row=4, column=c).border = border_all
        ws.cell(row=4, column=c).fill = fill_blue_hdr

    # ── Main Header Content (Row 5 — SEJAJAR!) ──
    # Columns A, B, C: BULAN, HARI, TANGGAL
    c_bln_hdr = ws.cell(row=5, column=1, value="BULAN")
    c_bln_hdr.font = f_header; c_bln_hdr.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    c_bln_hdr.border = border_all; c_bln_hdr.fill = fill_white

    c_hari_hdr = ws.cell(row=5, column=2, value="HARI")
    c_hari_hdr.font = f_header; c_hari_hdr.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    c_hari_hdr.border = border_all; c_hari_hdr.fill = fill_white

    c_tgl_hdr = ws.cell(row=5, column=3, value="TANGGAL")
    c_tgl_hdr.font = f_header; c_tgl_hdr.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    c_tgl_hdr.border = border_all; c_tgl_hdr.fill = fill_white

    # Columns p_start_col..p_end_col: Personnel names (vertical text 90°)
    for col_idx, p in enumerate(personnel_list, p_start_col):
        cell = ws.cell(row=5, column=col_idx, value=p["nama"])
        cell.font = f_header
        cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True)
        cell.border = border_all
        if p.get("is_resor"):
            cell.fill = fill_blue_hdr
        else:
            cell.fill = fill_white

    # Columns right_start_col..right_end_col: Signature Box (Kepala Urusan Preventif)
    kaur_prev = next((p for p in personnel_list if "Preventif" in p.get("jabatan", "")), None)
    if not kaur_prev:
        kaur_prev = next((p for p in personnel_list if p.get("role") == "KAUR"), None)
    k_prev_name = kaur_prev["nama"] if kaur_prev else "AGUS PRIYONO"
    k_prev_nipp = kaur_prev.get("nipp", "-") if kaur_prev else "-"
    nipp_line = f"\nNIPP. {k_prev_nipp}" if k_prev_nipp and k_prev_nipp != "-" else ""

    ws.merge_cells(start_row=5, start_column=right_start_col, end_row=5, end_column=right_end_col)
    c_sign = ws.cell(row=5, column=right_start_col, value=f"Kepala Urusan Preventif\n\n\n\n{k_prev_name}{nipp_line}")
    c_sign.font = Font(name=font_main, size=8.5, color="000000")
    c_sign.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c_sign.fill = fill_white
    for c in range(right_start_col, right_end_col + 1):
        ws.cell(row=5, column=c).border = border_all
        ws.cell(row=5, column=c).fill = fill_white

    # ── Row 6: Spacer Row (Left) & KELOMPOK PERALATAN (Right) ──
    for c in range(1, gap_col + 1):
        ws.cell(row=6, column=c).border = border_all
        ws.cell(row=6, column=c).fill = fill_white

    ws.merge_cells(start_row=6, start_column=right_start_col, end_row=6, end_column=right_end_col)
    c_kp = ws.cell(row=6, column=right_start_col, value="KELOMPOK PERALATAN")
    c_kp.font = f_header
    c_kp.alignment = al_center
    c_kp.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    for c in range(right_start_col, right_end_col + 1):
        ws.cell(row=6, column=c).border = border_all

    # Set row heights for headers
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 68
    ws.row_dimensions[6].height = 18.5

    # ── Populate Left Matrix Data (Rows 7 to 7 + num_days - 1) ──
    row_cursor = 7
    for day in range(1, num_days + 1):
        dt_obj = date(year, month, day)
        day_name = DAY_NAMES_ID[dt_obj.weekday()]
        is_sun = (dt_obj.weekday() == 6)

        ws.row_dimensions[row_cursor].height = 18.5

        ws.cell(row=row_cursor, column=1, value="")
        ws.cell(row=row_cursor, column=1).border = border_all

        c_hari = ws.cell(row=row_cursor, column=2, value=day_name)
        c_hari.font = f_cell
        c_hari.alignment = al_center
        c_hari.border = border_all

        c_tgl = ws.cell(row=row_cursor, column=3, value=day)
        c_tgl.font = f_cell_bold
        c_tgl.alignment = al_center
        c_tgl.border = border_all

        day_vals = matrix_data[day]
        for col_idx, p in enumerate(personnel_list, p_start_col):
            val = day_vals.get(p["key"], "P")
            cell = ws.cell(row=row_cursor, column=col_idx, value=val)
            cell.font = f_cell_bold if val != "P" else f_cell
            cell.alignment = al_center
            cell.border = border_all

            if is_sun:
                cell.fill = fill_sunday
                c_hari.fill = fill_sunday
                c_tgl.fill = fill_sunday
            else:
                cell.fill = fill_white

        row_cursor += 1

    # Merge Bulan column vertically
    ws.merge_cells(f"A7:A{row_cursor - 1}")
    b_cell = ws["A7"]
    b_cell.value = MONTH_NAMES_ID[month]
    b_cell.font = f_header
    b_cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)

    # Slogan below the matrix
    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=p_end_col)
    slogan_cell = ws.cell(row=row_cursor, column=1, value='"Utamakan Keselamatan, Berdo\'a sebelum bekerja"')
    slogan_cell.font = f_slogan
    slogan_cell.alignment = al_center
    ws.row_dimensions[row_cursor].height = 20
    row_cursor += 1

    # ── Right Side: 3-Column Legend (Kelompok Peralatan 1 s.d. 46) ──
    col_positions = [
        {"no_col": right_start_col, "desc_start": right_start_col + 1, "desc_end": right_start_col + 1},  # Block 1
        {"no_col": right_start_col + 2, "desc_start": right_start_col + 3, "desc_end": right_start_col + 4},  # Block 2
        {"no_col": right_start_col + 5, "desc_start": right_start_col + 6, "desc_end": right_start_col + 7},  # Block 3
    ]

    col_end_rows = [7, 7, 7]
    for c_idx, col_data in enumerate(LEGEND_COLUMNS):
        cur_r = 7
        no_c = col_positions[c_idx]["no_col"]
        ds_c = col_positions[c_idx]["desc_start"]
        de_c = col_positions[c_idx]["desc_end"]

        for station_name, group_ids in col_data:
            ws.merge_cells(start_row=cur_r, start_column=no_c, end_row=cur_r, end_column=de_c)
            st_cell = ws.cell(row=cur_r, column=no_c, value=station_name)
            st_cell.font = Font(name=font_main, size=8, bold=True, color="000000")
            st_cell.alignment = al_center
            st_color = STATION_COLORS.get(station_name, "D9EAD3")
            st_fill = PatternFill(start_color=st_color, end_color=st_color, fill_type="solid")
            st_cell.fill = st_fill
            for c in range(no_c, de_c + 1):
                ws.cell(row=cur_r, column=c).border = border_all
            ws.row_dimensions[cur_r].height = 20 if cur_r == 8 else 18.5
            cur_r += 1

            for gid in group_ids:
                g_info = TABLO_GROUPS.get(gid, {"name": ""})
                c_no = ws.cell(row=cur_r, column=no_c, value=gid)
                c_no.font = f_cell_bold
                c_no.alignment = al_center
                c_no.border = border_all
                c_no.fill = fill_white

                if ds_c == de_c:
                    c_desc = ws.cell(row=cur_r, column=ds_c, value=g_info["name"])
                    c_desc.font = Font(name=font_main, size=7.5, color="000000")
                    c_desc.alignment = al_left
                    c_desc.border = border_all
                    c_desc.fill = fill_white
                else:
                    ws.merge_cells(start_row=cur_r, start_column=ds_c, end_row=cur_r, end_column=de_c)
                    c_desc = ws.cell(row=cur_r, column=ds_c, value=g_info["name"])
                    c_desc.font = Font(name=font_main, size=7.5, color="000000")
                    c_desc.alignment = al_left
                    c_desc.fill = fill_white
                    for c in range(ds_c, de_c + 1):
                        ws.cell(row=cur_r, column=c).border = border_all

                ws.row_dimensions[cur_r].height = 18.5
                cur_r += 1

        col_end_rows[c_idx] = cur_r

    # ── Table Data Personil (Under Legend Col 2 & 3: cols dp_start..dp_end) ──
    dp_start = right_start_col + 2
    dp_end = right_end_col
    p_start = max(col_end_rows[1], col_end_rows[2]) + 2
    ws.merge_cells(start_row=p_start, start_column=dp_start, end_row=p_start, end_column=dp_end)
    p_hdr = ws.cell(row=p_start, column=dp_start, value="Data Personil")
    p_hdr.font = f_header
    p_hdr.alignment = al_center
    p_hdr.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    for c in range(dp_start, dp_end + 1): ws.cell(row=p_start, column=c).border = border_all
    ws.row_dimensions[p_start].height = 18.5

    p_sub_r = p_start + 1
    ws.row_dimensions[p_sub_r].height = 18.5
    # Nama: dp_start..dp_start + 1
    ws.merge_cells(start_row=p_sub_r, start_column=dp_start, end_row=p_sub_r, end_column=dp_start + 1)
    ws.cell(row=p_sub_r, column=dp_start, value="Nama").font = f_header
    ws.cell(row=p_sub_r, column=dp_start).alignment = al_center
    ws.cell(row=p_sub_r, column=dp_start).border = border_all
    ws.cell(row=p_sub_r, column=dp_start + 1).border = border_all

    # NIPP: dp_start + 2
    ws.cell(row=p_sub_r, column=dp_start + 2, value="NIPP").font = f_header
    ws.cell(row=p_sub_r, column=dp_start + 2).alignment = al_center
    ws.cell(row=p_sub_r, column=dp_start + 2).border = border_all

    # Jabatan: dp_start + 3..dp_start + 4
    ws.merge_cells(start_row=p_sub_r, start_column=dp_start + 3, end_row=p_sub_r, end_column=dp_start + 4)
    ws.cell(row=p_sub_r, column=dp_start + 3, value="Jabatan").font = f_header
    ws.cell(row=p_sub_r, column=dp_start + 3).alignment = al_center
    ws.cell(row=p_sub_r, column=dp_start + 3).border = border_all
    ws.cell(row=p_sub_r, column=dp_start + 4).border = border_all

    # Domisili: dp_end
    ws.cell(row=p_sub_r, column=dp_end, value="Domisili").font = f_header
    ws.cell(row=p_sub_r, column=dp_end).alignment = al_center
    ws.cell(row=p_sub_r, column=dp_end).border = border_all

    cur_pr = p_sub_r + 1
    for p in personnel_list:
        ws.row_dimensions[cur_pr].height = 18.5
        ws.merge_cells(start_row=cur_pr, start_column=dp_start, end_row=cur_pr, end_column=dp_start + 1)
        c1 = ws.cell(row=cur_pr, column=dp_start, value=p["full_name"])
        c1.font = Font(name=font_main, size=8, color="000000")
        c1.alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=cur_pr, column=dp_start).border = border_all
        ws.cell(row=cur_pr, column=dp_start + 1).border = border_all

        c2 = ws.cell(row=cur_pr, column=dp_start + 2, value=p["nipp"])
        c2.font = Font(name=font_main, size=8, color="000000")
        c2.alignment = al_center
        c2.border = border_all

        ws.merge_cells(start_row=cur_pr, start_column=dp_start + 3, end_row=cur_pr, end_column=dp_start + 4)
        c3 = ws.cell(row=cur_pr, column=dp_start + 3, value=p["jabatan"])
        c3.font = Font(name=font_main, size=7.5, color="000000")
        c3.alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=cur_pr, column=dp_start + 3).border = border_all
        ws.cell(row=cur_pr, column=dp_start + 4).border = border_all

        c4 = ws.cell(row=cur_pr, column=dp_end, value=p["domisili"])
        c4.font = Font(name=font_main, size=8, color="000000")
        c4.alignment = al_center
        c4.border = border_all
        cur_pr += 1

    # ── Approval Signature Block (Under Data Personil: cols dp_start..dp_end) ──
    sign_r = cur_pr + 2
    ws.merge_cells(start_row=sign_r, start_column=dp_start, end_row=sign_r, end_column=dp_end)
    s1 = ws.cell(row=sign_r, column=dp_start, value="Disetujui Oleh :")
    s1.font = f_cell; s1.alignment = al_center
    ws.row_dimensions[sign_r].height = 18.5

    sign_r += 1
    ws.merge_cells(start_row=sign_r, start_column=dp_start, end_row=sign_r, end_column=dp_end)
    s2 = ws.cell(row=sign_r, column=dp_start, value="KUPT Resor Sintelis 1.21 Bogor")
    s2.font = f_header; s2.alignment = al_center
    ws.row_dimensions[sign_r].height = 18.5

    for _ in range(3):
        sign_r += 1
        ws.row_dimensions[sign_r].height = 18.5

    sign_r += 1
    ws.merge_cells(start_row=sign_r, start_column=dp_start, end_row=sign_r, end_column=dp_end)
    s3 = ws.cell(row=sign_r, column=dp_start, value=resor_nama)
    s3.font = Font(name=font_main, size=9, bold=True, underline="single"); s3.alignment = al_center
    ws.row_dimensions[sign_r].height = 18.5

    sign_r += 1
    ws.merge_cells(start_row=sign_r, start_column=dp_start, end_row=sign_r, end_column=dp_end)
    s4 = ws.cell(row=sign_r, column=dp_start, value=f"NIPP. {resor_nipp}")
    s4.font = f_cell; s4.alignment = al_center
    ws.row_dimensions[sign_r].height = 18.5

    # ── Column Width Adjustments ──
    ws.column_dimensions["A"].width = 5.0   # Bulan
    ws.column_dimensions["B"].width = 9.5   # Hari
    ws.column_dimensions["C"].width = 5.5   # Tanggal
    for c in range(p_start_col, p_end_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 8.5
    ws.column_dimensions[get_column_letter(gap_col)].width = 2.0
    ws.column_dimensions[get_column_letter(right_start_col)].width = 4.5
    ws.column_dimensions[get_column_letter(right_start_col + 1)].width = 32.0
    ws.column_dimensions[get_column_letter(right_start_col + 2)].width = 4.5
    ws.column_dimensions[get_column_letter(right_start_col + 3)].width = 20.0
    ws.column_dimensions[get_column_letter(right_start_col + 4)].width = 11.0
    ws.column_dimensions[get_column_letter(right_start_col + 5)].width = 4.5
    ws.column_dimensions[get_column_letter(right_start_col + 6)].width = 20.0
    ws.column_dimensions[get_column_letter(right_start_col + 7)].width = 11.0

    max_print_row = max(row_cursor, sign_r)
    ws.print_area = f"A1:{get_column_letter(right_end_col)}{max_print_row}"

    month_name = MONTH_NAMES_ID.get(month, f"BULAN_{month}")
    tablo_filename = f"TABLO {month_name} {year}.xlsx"

    if output_path:
        out_p = Path(output_path)
        if out_p.is_dir() or str(output_path).endswith(("\\", "/")):
            output_path = out_p / tablo_filename
        elif "tablo_perawatan_berkala" in out_p.name.lower() or "tablo_kustom" in out_p.name.lower():
            output_path = out_p.parent / tablo_filename
        else:
            output_path = out_p
    else:
        output_path = APP_DIR / "logs" / tablo_filename

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        try:
            output_path.unlink()
        except Exception:
            pass

    final_out = output_path.parent.resolve() / output_path.name
    try:
        wb.save(output_path)
        print(f"[OUTPUT_FILE] {final_out}")
        print(f"[OK] Tablo Excel successfully generated: {final_out}")
        return str(final_out)
    except PermissionError:
        fallback = output_path.parent.resolve() / f"{output_path.stem}_updated.xlsx"
        wb.save(fallback)
        print(f"[OUTPUT_FILE] {fallback}")
        print(f"[WARN] File locked. Saved to fallback: {fallback}")
        return str(fallback)


def main():
    parser = argparse.ArgumentParser(description="Export official KAI Tablo Perawatan Excel.")
    parser.add_argument("--mode", choices=["pipeline", "custom"], default="pipeline", help="Export mode: 'pipeline' (schedule.json) or 'custom' (PDF folder)")
    parser.add_argument("--folder", default=None, help="Custom PDF target folder for mode=custom")
    parser.add_argument("--schedule", default=str(APP_DIR / "schedule.json"), help="Path to schedule.json")
    parser.add_argument("--output", default=None, help="Path to output .xlsx (default: logs/TABLO [BULAN] [TAHUN].xlsx)")
    parser.add_argument("--output-dir", default=None, help="Directory to save TABLO [BULAN] [TAHUN].xlsx")
    parser.add_argument("--year", type=int, default=None, help="Specific year to export (optional)")
    parser.add_argument("--month", type=int, default=None, help="Specific month 1..12 to export (optional)")
    args = parser.parse_args()

    if args.output:
        output_path = Path(args.output).resolve()
    elif args.output_dir:
        output_path = Path(args.output_dir).resolve()
    else:
        output_path = None

    if args.mode == "custom":
        if not args.folder or not Path(args.folder).exists():
            print(f"[ERROR] Folder kustom '{args.folder}' tidak ditemukan.", file=sys.stderr)
            return 1
        custom_pdf_dir = Path(args.folder).resolve()
        temp_sched_file = APP_DIR / "logs" / "temp_custom_schedule.json"
        temp_sched_file.parent.mkdir(parents=True, exist_ok=True)

        print(f"⚡ [TABLO MODE 2] Mengekstrak jadwal langsung dari folder: {custom_pdf_dir}...")
        try:
            from scheduler import build_schedule, load_mapping, load_data_acuan
        except ImportError:
            sys.path.insert(0, str(Path(__file__).resolve().parent))
            from scheduler import build_schedule, load_mapping, load_data_acuan

        engine_dir = Path(__file__).resolve().parent
        m_path = engine_dir / "asset_waktu_mapping.json"
        if not m_path.exists() and (APP_DIR / "config" / "asset_waktu_mapping.json").exists():
            m_path = APP_DIR / "config" / "asset_waktu_mapping.json"
        mapping = load_mapping(m_path) if m_path.exists() else {}

        a_path = engine_dir / "data_acuan_tenaga_gabungan.json"
        if not a_path.exists() and (APP_DIR / "config" / "data_acuan_tenaga_gabungan.json").exists():
            a_path = APP_DIR / "config" / "data_acuan_tenaga_gabungan.json"
        acuan = load_data_acuan(a_path) if a_path.exists() else {}

        photos_dir = APP_DIR / "03_photos_export"
        sched_data = build_schedule(
            pdf_dir=custom_pdf_dir,
            photos_dir=photos_dir,
            mapping=mapping,
            acuan=acuan,
            jam_mulai=7 * 60,
            jam_selesai=18 * 60,
            tim_max=2
        )
        try:
            with open(temp_sched_file, "w", encoding="utf-8") as f:
                json.dump(sched_data, f, indent=2, ensure_ascii=False)
        except Exception:
            pass
    else:
        # Mode pipeline
        sched_file = Path(args.schedule).resolve()
        if not sched_file.exists():
            print(f"[ERROR] Berkas jadwal {sched_file} tidak ditemukan. Silakan jalankan Step 3 Scheduler terlebih dahulu.", file=sys.stderr)
            return 1
        with open(sched_file, "r", encoding="utf-8") as f:
            sched_data = json.load(f)

    generate_tablo_workbook(sched_data, year=args.year, month=args.month, output_path=output_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
